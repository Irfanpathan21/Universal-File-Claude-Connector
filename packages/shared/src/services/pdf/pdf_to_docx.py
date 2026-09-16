import sys
import os
import io

def convert_pdf_to_docx(input_pdf, output_docx):
    # Method 1: Try pdf2docx (high fidelity document reconstruction)
    try:
        from pdf2docx import Converter
        cv = Converter(input_pdf)
        cv.convert(output_docx, start=0, end=None)
        cv.close()
        if os.path.exists(output_docx) and os.path.getsize(output_docx) > 0:
            return True, "Converted using pdf2docx"
    except Exception as e:
        sys.stderr.write(f"pdf2docx notice: {e}\n")

    # Method 2: Fallback to PyMuPDF (fitz) + python-docx
    try:
        import fitz
        from docx import Document
        
        doc = fitz.open(input_pdf)
        docx_doc = Document()
        has_content = False
        
        for pno in range(len(doc)):
            page = doc[pno]
            text = page.get_text("text")
            if text and text.strip():
                has_content = True
                paragraphs = text.split("\n")
                for para_text in paragraphs:
                    if para_text.strip():
                        docx_doc.add_paragraph(para_text)
            
            if pno < len(doc) - 1:
                docx_doc.add_page_break()
                
        docx_doc.save(output_docx)
        doc.close()
        if os.path.exists(output_docx) and os.path.getsize(output_docx) > 0:
            return True, "Converted using fitz + python-docx fallback"
    except Exception as e:
        sys.stderr.write(f"fitz + python-docx notice: {e}\n")

    # Method 3: Fallback to pypdf + docx
    try:
        from pypdf import PdfReader
        from docx import Document
        
        reader = PdfReader(input_pdf)
        docx_doc = Document()
        for idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                for line in text.split("\n"):
                    if line.strip():
                        docx_doc.add_paragraph(line)
            if idx < len(reader.pages) - 1:
                docx_doc.add_page_break()
                
        docx_doc.save(output_docx)
        if os.path.exists(output_docx) and os.path.getsize(output_docx) > 0:
            return True, "Converted using pypdf fallback"
    except Exception as e:
        sys.stderr.write(f"pypdf notice: {e}\n")

    return False, "All conversion methods failed"

def main():
    if len(sys.argv) < 3:
        sys.stderr.write("Usage: python pdf_to_docx.py <input_pdf> <output_docx>\n")
        sys.exit(1)
        
    input_pdf = os.path.abspath(sys.argv[1])
    output_docx = os.path.abspath(sys.argv[2])
    
    if not os.path.exists(input_pdf):
        sys.stderr.write(f"Input file does not exist: {input_pdf}\n")
        sys.exit(1)
        
    os.makedirs(os.path.dirname(output_docx), exist_ok=True)
    
    success, msg = convert_pdf_to_docx(input_pdf, output_docx)
    if success:
        sys.stdout.write(f"SUCCESS: {msg}\n")
        sys.exit(0)
    else:
        sys.stderr.write(f"FAILURE: {msg}\n")
        sys.exit(1)

if __name__ == '__main__':
    main()
